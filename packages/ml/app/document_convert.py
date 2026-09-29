"""Приём DOCX и XML: преобразование в PDF на входе (ТЗ 9.1).

Весь конвейер — страницы, цитаты, координаты доказательств — построен на
PDF. Вторая ветка разбора для каждого формата разошлась бы с первой, поэтому
DOCX и XML приводятся к PDF один раз, при загрузке, а исходный формат и его
SHA-256 сохраняются в метаданных документа: доказательство ссылается на
страницу преобразованного PDF, и это видно, а не скрыто.

Ограничения честно названы. Из DOCX переносится текст абзацев и таблиц, без
вёрстки и графики; из XML — текст элементов с путём к ним. Для
текстовых документов (акты, ведомости, журналы) этого достаточно для
сравнения значений; чертежи в DOCX/XML не приходят.
"""
from __future__ import annotations

import hashlib
import io
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass

FORMAT_PDF, FORMAT_DOCX, FORMAT_XML = "PDF", "DOCX", "XML"


class UnsupportedFormatError(ValueError):
    """Файл не PDF, не DOCX и не XML — или повреждён так, что им не является."""


@dataclass
class Converted:
    pdf: bytes
    source_format: str
    source_sha256: str


def detect(data: bytes) -> str | None:
    """Формат по содержимому, а не по расширению (Б.4)."""
    if data.startswith(b"%PDF-"):
        return FORMAT_PDF
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if "word/document.xml" in archive.namelist():
                    return FORMAT_DOCX
        except zipfile.BadZipFile:
            return None
        return None
    head = data.lstrip(b"\xef\xbb\xbf \t\r\n")[:5]
    return FORMAT_XML if head.startswith((b"<?xml", b"<")) else None


def _docx_lines(data: bytes) -> list[str]:
    from docx import Document

    document = Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return lines


def _xml_lines(data: bytes) -> list[str]:
    # DTD запрещён целиком: через него идут и внешние сущности (XXE), и
    # раздувание сущностей. Документам ИД объявление DTD не нужно.
    if b"<!DOCTYPE" in data[:4096].upper() or b"<!ENTITY" in data.upper():
        raise UnsupportedFormatError("XML с объявлением DTD/сущностей не принимается")
    try:
        root = ET.fromstring(data)  # noqa: S314 — DTD отклонён выше: без него XXE и раздувания нет
    except ET.ParseError as exc:
        raise UnsupportedFormatError(f"XML не читается: {exc}") from exc
    lines: list[str] = []

    def walk(node: ET.Element, path: str) -> None:
        tag = node.tag.split("}")[-1]
        here = f"{path}/{tag}" if path else tag
        attrs = " ".join(f"{k.split('}')[-1]}={v}" for k, v in node.attrib.items())
        text = (node.text or "").strip()
        if text or attrs:
            lines.append(f"{here}: {text} {attrs}".strip())
        for child in node:
            walk(child, here)

    walk(root, "")
    return lines


_FONT_CANDIDATES = ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                    "/usr/share/fonts/dejavu/DejaVuSans.ttf")


def _font() -> str:
    """Шрифт с кириллицей для reportlab; без него — встроенный (латиница)."""
    from pathlib import Path

    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if "DejaVuSans" in pdfmetrics.getRegisteredFontNames():
        return "DejaVuSans"
    path = next((item for item in _FONT_CANDIDATES if Path(item).is_file()), None)
    if path is None:
        return "Helvetica"
    pdfmetrics.registerFont(TTFont("DejaVuSans", path))
    return "DejaVuSans"


def _to_pdf(lines: list[str], title: str) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    font = _font()
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = font
    stream = io.BytesIO()
    story = [Paragraph(escape(title), styles["Heading2"])]
    story += [Paragraph(escape(line), styles["BodyText"]) for line in lines] or [
        Paragraph("Текста в документе нет.", styles["BodyText"])]
    SimpleDocTemplate(stream, pagesize=A4, title=title).build(story)
    return stream.getvalue()


def to_pdf(data: bytes, filename: str) -> Converted:
    """PDF — как есть; DOCX и XML — в PDF. Иное — UnsupportedFormatError."""
    source_format = detect(data)
    digest = hashlib.sha256(data).hexdigest()
    if source_format == FORMAT_PDF:
        return Converted(data, FORMAT_PDF, digest)
    if source_format is None:
        raise UnsupportedFormatError("неподдерживаемый формат; поддерживаются PDF, DOCX, XML")
    try:
        lines = _docx_lines(data) if source_format == FORMAT_DOCX else _xml_lines(data)
    except UnsupportedFormatError:
        raise
    except Exception as exc:  # noqa: BLE001 — повреждённый файл: отказ с причиной
        raise UnsupportedFormatError(f"{source_format} повреждён: {type(exc).__name__}") from exc
    return Converted(_to_pdf(lines, filename), source_format, digest)

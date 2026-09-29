"""Приём DOCX и XML (ТЗ 9.1): преобразование в PDF, отказ с причиной.

Проверяется главное: текст исходного документа доезжает до PDF и читается
тем же разбором, что и обычный PDF, — иначе приём формата был бы видимостью.
"""
import hashlib
import io
import sys
from pathlib import Path

import pymupdf
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import document_convert  # noqa: E402


def _docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph("Класс бетона В30 по проекту")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Параметр"
    table.rows[0].cells[1].text = "Значение 120 м²"
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


XML = "<?xml version='1.0' encoding='utf-8'?><act><number>42</number>" \
      "<work>Армирование плиты</work></act>".encode()


def _text(pdf: bytes) -> str:
    with pymupdf.open(stream=pdf, filetype="pdf") as doc:
        return " ".join(page.get_text() for page in doc)


def test_formats_are_detected_by_content_not_extension():
    assert document_convert.detect(b"%PDF-1.7 ...") == "PDF"
    assert document_convert.detect(_docx()) == "DOCX"
    assert document_convert.detect(XML) == "XML"
    assert document_convert.detect(b"PK\x03\x04not-a-docx") is None
    assert document_convert.detect(b"\x89PNG....") is None


def test_docx_text_and_tables_reach_the_pdf():
    data = _docx()
    converted = document_convert.to_pdf(data, "акт.docx")
    text = _text(converted.pdf)
    assert "Класс бетона В30" in text and "120 м²" in text
    assert converted.source_format == "DOCX"
    assert converted.source_sha256 == hashlib.sha256(data).hexdigest()


def test_xml_elements_reach_the_pdf_with_their_path():
    converted = document_convert.to_pdf(XML, "акт.xml")
    text = _text(converted.pdf)
    assert "act/work: Армирование плиты" in text
    assert converted.source_sha256 == hashlib.sha256(XML).hexdigest()


def test_xml_with_dtd_is_refused():
    """XXE и раздувание сущностей идут через DTD — такой XML не принимается."""
    bomb = b'<?xml version="1.0"?><!DOCTYPE a [<!ENTITY x "xx">]><a>&x;</a>'
    with pytest.raises(document_convert.UnsupportedFormatError, match="DTD"):
        document_convert.to_pdf(bomb, "a.xml")


def test_broken_docx_is_refused_with_a_reason():
    with pytest.raises(document_convert.UnsupportedFormatError, match="неподдерживаемый"):
        document_convert.to_pdf(b"PK\x03\x04broken", "x.docx")
